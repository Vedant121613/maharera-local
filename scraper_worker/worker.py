import os
import sys
import time
import sqlite3
import threading
import subprocess
from pathlib import Path

from dotenv import load_dotenv
load_dotenv()

import api_client as api
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(exist_ok=True)

LINK_SCRAPER = BASE_DIR / "link_scraper.py"
DATA_SCRAPER = BASE_DIR / "data_scraper.py"

POLL_INTERVAL = int(os.getenv("POLL_INTERVAL", "5"))
STATUS_CHECK_INTERVAL = int(os.getenv("STATUS_CHECK_INTERVAL", "3"))  # was 10s — Stop felt slow to react


def safe_name(district: str) -> str:
    return "".join(c if c.isalnum() else "_" for c in district)


def run_subprocess_with_stop_watch(cmd, cwd, env, job_id):
    """Runs cmd, polling the backend for STOP_REQUESTED and touching a stop
    flag file (already understood by link_scraper.py / data_scraper.py) when
    seen. Returns (returncode, was_stopped)."""
    proc = subprocess.Popen(cmd, cwd=str(cwd), env=env)
    stop_flag = Path(env["STOP_FLAG_PATH"])
    was_stopped = False
    last_check = 0.0

    while proc.poll() is None:
        time.sleep(1)
        if time.time() - last_check >= STATUS_CHECK_INTERVAL:
            last_check = time.time()
            try:
                job = api.get_job(job_id)
                if job and job.get("status") == "STOP_REQUESTED" and not stop_flag.exists():
                    print(f"[worker] STOP_REQUESTED for job {job_id}; signalling scraper...")
                    stop_flag.touch()
                    was_stopped = True
            except Exception as e:
                print(f"[worker] status check failed: {e}")

    proc.wait()
    if stop_flag.exists():
        was_stopped = True
    return proc.returncode, was_stopped


def read_links_xlsx(path: Path):
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        rows.append({
            "reraId": rec.get("RERA ID"),
            "projectName": rec.get("Project Name"),
            "pincode": rec.get("Pincode"),
            "viewUrl": rec.get("View Details URL"),
            "pageNumber": rec.get("Page Number"),
            "certificateUrl": rec.get("Certificate URL"),
            "certificateStatus": rec.get("Certificate Status"),
        })
    wb.close()
    return rows



def run_link_job(job):
    job_id, district, district_id = job["id"], job["district"], job["districtId"]
    print(f"[worker] Link job {job_id}: {district} (id={district_id})")

    output_file = DATA_DIR / f"links_{safe_name(district)}.xlsx"
    stop_flag = DATA_DIR / f"stop_link_{job_id}.flag"
    if stop_flag.exists():
        stop_flag.unlink()

    env = os.environ.copy()
    env.update({
        "SCRAPE_STATE_ID": "27",
        "SCRAPE_DISTRICT_ID": str(district_id),
        "SCRAPE_DISTRICT_NAME": district,
        "SCRAPE_OUTPUT_FILE": str(output_file),
        "STOP_FLAG_PATH": str(stop_flag),
    })

    returncode, was_stopped = run_subprocess_with_stop_watch(
        [sys.executable, str(LINK_SCRAPER)], DATA_DIR, env, job_id
    )

    rows = read_links_xlsx(output_file)
    inserted = api.post_links(district, rows) if rows else 0
    api.post_progress(job_id, processed=len(rows), total=len(rows), found=inserted)

    if was_stopped:
        api.post_complete(job_id, "STOPPED")
    elif returncode == 0:
        api.post_complete(job_id, "COMPLETED")
    else:
        api.post_complete(job_id, "FAILED", error=f"link_scraper.py exited with code {returncode}")

    if stop_flag.exists():
        stop_flag.unlink()
    print(f"[worker] Link job {job_id} done: {len(rows)} links, {inserted} newly inserted.")


def build_input_excel(district, path: Path):
    rows = api.get_links_for_district(district)
    wb = Workbook()
    ws = wb.active
    ws.append(["RERA ID", "View Details URL", "Page Number"])
    for r in rows:
        ws.append([r.get("reraId"), r.get("viewUrl"), r.get("pageNumber")])
    wb.save(path)
    return len(rows)


def read_basic_data_sqlite(db_file: Path):
    if not db_file.exists():
        return []
    conn = sqlite3.connect(str(db_file))
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()
    cur.execute("SELECT * FROM projects")
    rows = [dict(r) for r in cur.fetchall()]
    conn.close()
    out = []
    for r in rows:
        out.append({
            "projectId": r.get("project_id"),
            "reraId": r.get("rera_id"),
            "district": r.get("district"),
            "taluka": r.get("taluka"),
            "village": r.get("village"),
            "projectName": r.get("project_name"),
            "registrationNumber": r.get("registration_number"),
            "dateOfRegistration": r.get("date_of_registration"),
            "proposedCompletionDate": r.get("proposed_completion_date_original"),
            "address": r.get("address"),
            "pincode": r.get("pin_code"),
            "longitude": r.get("longitude"),
            "latitude": r.get("latitude"),
            "viewDetailsUrl": r.get("view_details_url"),
            "pageNumber": r.get("page_number"),
            "status": r.get("status"),
        })
    return out


def run_data_job(job):
    job_id, district, district_id = job["id"], job["district"], job["districtId"]
    print(f"[worker] Data job {job_id}: {district} (id={district_id})")

    input_file = DATA_DIR / f"links_{safe_name(district)}.xlsx"
    total_links = build_input_excel(district, input_file)
    if total_links == 0:
        api.post_complete(job_id, "FAILED", error="No links found for this district. Run the Link Scraper first.")
        print(f"[worker] Data job {job_id} skipped: no links for {district}.")
        return

    output_file = DATA_DIR / f"basicdata_{safe_name(district)}.xlsx"
    failures_file = DATA_DIR / f"basicdata_{safe_name(district)}_failures.xlsx"
    db_file = DATA_DIR / f"basicdata_{safe_name(district)}.db"
    stop_flag = DATA_DIR / f"stop_data_{job_id}.flag"
    if stop_flag.exists():
        stop_flag.unlink()

    env = os.environ.copy()
    env.update({
        "SCRAPE_DISTRICT_NAME": district,
        "SCRAPE_INPUT_FILE": str(input_file),
        "SCRAPE_OUTPUT_FILE": str(output_file),
        "SCRAPE_FAILURES_FILE": str(failures_file),
        "SCRAPE_DB_FILE": str(db_file),
        "STOP_FLAG_PATH": str(stop_flag),
    })

    api.post_progress(job_id, total=total_links)

    returncode, was_stopped = run_subprocess_with_stop_watch(
        [sys.executable, str(DATA_SCRAPER)], DATA_DIR, env, job_id
    )

    rows = read_basic_data_sqlite(db_file)
    if rows:
        api.post_basic_data(rows)
    succeeded = sum(1 for r in rows if r["status"] == "SUCCESS")
    failed = sum(1 for r in rows if r["status"] == "FAILED")
    api.post_progress(job_id, processed=len(rows), total=total_links, found=succeeded, failed=failed)

    if was_stopped:
        api.post_complete(job_id, "STOPPED")
    elif returncode == 0:
        api.post_complete(job_id, "COMPLETED")
    else:
        api.post_complete(job_id, "FAILED", error=f"data_scraper.py exited with code {returncode}")

    if stop_flag.exists():
        stop_flag.unlink()
    print(f"[worker] Data job {job_id} done: {len(rows)} processed, {succeeded} success, {failed} failed.")


STOP_EVENT = threading.Event()


def poll_loop(job_type, run_fn):
    """Independent poll loop for one job type ('link' or 'data'). Runs on its
    own thread so a long Link job can never delay or block a Data job (or
    vice versa) — each has its own Start/Stop lifecycle end to end."""
    print(f"[worker:{job_type}] loop started")
    while not STOP_EVENT.is_set():
        try:
            job = api.get_next_job(job_type)
            if not job:
                time.sleep(POLL_INTERVAL)
                continue
            run_fn(job)
        except Exception as e:
            print(f"[worker:{job_type}] Unhandled error: {e}. Retrying in {POLL_INTERVAL}s...")
            time.sleep(POLL_INTERVAL)


def heartbeat_loop():
    while not STOP_EVENT.is_set():
        try:
            api.heartbeat()
        except Exception as e:
            print(f"[worker] heartbeat failed: {e}")
        time.sleep(POLL_INTERVAL)


def main():
    print("=" * 60)
    print("MahaRERA scraper worker starting")
    print(f"API_BASE_URL = {os.getenv('API_BASE_URL')}")
    print("=" * 60)

    threads = [
        threading.Thread(target=heartbeat_loop, daemon=True),
        threading.Thread(target=poll_loop, args=("link", run_link_job), daemon=True),
        threading.Thread(target=poll_loop, args=("data", run_data_job), daemon=True),
    ]
    for t in threads:
        t.start()

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        print("\n[worker] Shutting down (Ctrl+C)...")
        STOP_EVENT.set()


if __name__ == "__main__":
    main()